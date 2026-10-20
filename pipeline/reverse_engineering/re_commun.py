"""
SXMXDX - Reverse engineering : outils communs.

Ce module rend la chaîne indépendante d'un classeur particulier :

- dossier de sortie et configuration choisis à l'exécution (plus de chemin figé) ;
- détection AUTOMATIQUE de la mise en page temporelle de chaque feuille
  (mois en colonnes ou en lignes, colonne/ligne de total annuel) à partir des en-têtes,
  au lieu de supposer « Synthese, mois en C:N, total en O » ;
- helpers pour les particularités Excel (formules matricielles, feuilles graphiques).
"""
import json
import os
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl.utils import column_index_from_string, get_column_letter

DOSSIER_SORTIE_DEFAUT = Path("output/reverse_engineering")
EXTENSIONS_EXCEL = {".xlsx", ".xlsm", ".xltx", ".xltm"}


# ============================================================
# ARGUMENTS COMMUNS : --sortie DOSSIER   --config FICHIER.json
# ============================================================

def lire_arguments():
    """
    Retire les options communes de sys.argv, les place dans l'environnement
    (pour que tous les scripts de la chaîne les partagent) et renvoie les
    arguments positionnels restants (sans le nom du script).
    """
    restant, i = [], 1
    while i < len(sys.argv):
        arg = sys.argv[i]
        if arg in ("--sortie", "--config") and i + 1 < len(sys.argv):
            os.environ["RE_SORTIE" if arg == "--sortie" else "RE_CONFIG"] = sys.argv[i + 1]
            i += 2
            continue
        restant.append(arg)
        i += 1
    return restant


def dossier_sortie():
    return Path(os.environ.get("RE_SORTIE") or DOSSIER_SORTIE_DEFAUT)


def verifier_fichier_excel(chemin):
    fichier = Path(chemin)
    if not fichier.exists():
        raise FileNotFoundError(f"Fichier introuvable : {fichier}")
    if fichier.suffix.lower() not in EXTENSIONS_EXCEL:
        raise ValueError(
            f"Format non pris en charge : {fichier.suffix or '(sans extension)'}. "
            f"Formats acceptés : {', '.join(sorted(EXTENSIONS_EXCEL))}. "
            "Pour un .xls ou .csv, enregistrez d'abord le fichier en .xlsx."
        )
    return fichier


# ============================================================
# CONFIGURATION (facultative)
# ============================================================

CONFIG_DEFAUT = {
    # {"enero": 1, "januar": 1, ...} : noms de mois d'autres langues
    "mois_supplementaires": {},
    # libellés à ne jamais considérer comme concept métier (unités, devises...)
    "valeurs_non_semantiques": [],
    # {"VENTES": ["sales", "vente"], ...} : mots-clés de feuille -> domaine
    "domaines": {},
    # "auto" | "aucun" | ["NomFeuille", ...]
    "regroupement_mensuel": "auto",
    # forçage de la mise en page : {"Feuille": {"orientation": "COLONNES",
    #                                  "mois": "C:N", "annuelle": "O"}}
    "mise_en_page": {},
    # regroupe les formules recopiees (I2=F2+G2, I3=F3+G3...) en une seule regle
    "regroupement_formules_recopiees": True,
    "nom_modele": None,
    "version": None,
    "distance_max_libelle": 10,
    "mapping_metier": None,
}

_cache_config = {}


def charger_config():
    """Ordre de recherche : --config / RE_CONFIG, ./re_config.json, <sortie>/re_config.json."""
    candidats = []
    if os.environ.get("RE_CONFIG"):
        candidats.append(Path(os.environ["RE_CONFIG"]))
    candidats += [Path("re_config.json"), dossier_sortie() / "re_config.json"]

    for chemin in candidats:
        if chemin.exists():
            cle = str(chemin.resolve())
            if cle not in _cache_config:
                with chemin.open("r", encoding="utf-8") as f:
                    brut = json.load(f)
                _cache_config[cle] = {**CONFIG_DEFAUT, **brut}
            return _cache_config[cle]
    return dict(CONFIG_DEFAUT)


# ============================================================
# TEXTE
# ============================================================

def sans_accents(valeur):
    texte = unicodedata.normalize("NFD", str(valeur))
    return "".join(c for c in texte if unicodedata.category(c) != "Mn")


def normaliser_nom(valeur):
    """'Charges salariales' -> 'CHARGES_SALARIALES'"""
    texte = re.sub(r"[^A-Z0-9]+", "_", sans_accents(valeur or "").strip().upper())
    return texte.strip("_")


# ============================================================
# MOIS
# ============================================================

_MOIS_BASE = {
    1: ["janvier", "janv", "jan", "january"],
    2: ["fevrier", "fevr", "fev", "feb", "february"],
    3: ["mars", "mar", "march"],
    4: ["avril", "avr", "apr", "april"],
    5: ["mai", "may"],
    6: ["juin", "jun", "june"],
    7: ["juillet", "juil", "jul", "july"],
    8: ["aout", "aou", "aug", "august"],
    9: ["septembre", "sept", "sep", "september"],
    10: ["octobre", "oct", "october"],
    11: ["novembre", "nov", "november"],
    12: ["decembre", "dec", "december"],
}

_MOT_VERS_MOIS = None


def _table_mois():
    global _MOT_VERS_MOIS
    if _MOT_VERS_MOIS is None:
        table = {mot: num for num, mots in _MOIS_BASE.items() for mot in mots}
        for mot, num in (charger_config().get("mois_supplementaires") or {}).items():
            table[sans_accents(mot).lower().strip(". ")] = int(num)
        _MOT_VERS_MOIS = table
    return _MOT_VERS_MOIS


def tous_les_noms_de_mois():
    return set(_table_mois())


def numero_mois(valeur):
    """'Janvier', 'jan-25', 'Janv.' ou 'January 2025' -> 1 ; sinon None."""
    if not isinstance(valeur, str):
        return None
    texte = sans_accents(valeur).lower().strip()
    m = re.fullmatch(r"([a-z]+)\.?(?:[\s\-_/']*\d{2,4})?", texte)
    return _table_mois().get(m.group(1)) if m else None


# ============================================================
# FORMULES EXCEL
# ============================================================

def texte_formule(valeur):
    """
    Renvoie le texte d'une formule, qu'elle soit simple ('=A1+B1') ou matricielle
    (objet ArrayFormula d'openpyxl), sinon None.
    """
    if isinstance(valeur, str) and valeur.startswith("="):
        return valeur
    texte = getattr(valeur, "text", None)  # ArrayFormula
    if isinstance(texte, str) and texte.startswith("="):
        return texte
    return None


def feuilles_lisibles(workbook):
    """Feuilles de calcul (les feuilles-graphiques n'ont ni cellules ni formules)."""
    return [ws for ws in workbook.worksheets if hasattr(ws, "iter_rows")]


# ============================================================
# MISE EN PAGE TEMPORELLE
# ============================================================

MOTS_TOTAL = ("total", "annuel", "annee", "annual", "year", "cumul", "ytd")


def _plus_longue_suite(positions_mois):
    """
    positions_mois : [(position, numero_mois)] triés par position.
    Retourne la plus longue suite de mois consécutifs (1,2,3... avec retour 12->1).
    """
    meilleure, courante = [], []
    for position, mois in positions_mois:
        if courante and mois == courante[-1][1] % 12 + 1:
            courante.append((position, mois))
        else:
            if len(courante) > len(meilleure):
                meilleure = courante
            courante = [(position, mois)]
    return courante if len(courante) > len(meilleure) else meilleure


def detecter_mise_en_page(ws, min_mois=6):
    """
    Cherche dans la feuille une ligne (ou une colonne) d'en-têtes de mois.

    Retourne None si rien de convaincant, sinon :
    {
      "orientation": "COLONNES" | "LIGNES",
      "axe_en_tete": numéro de la ligne (COLONNES) ou de la colonne (LIGNES) d'en-tête,
      "positions": {numéro_colonne_ou_ligne: numéro_mois},
      "annuelle": numéro de la colonne/ligne de total, ou None,
      "origine": "DETECTEE"
    }
    """
    if not hasattr(ws, "iter_rows"):
        return None

    texte_ligne, texte_col = defaultdict(list), defaultdict(list)
    date_ligne, date_col = defaultdict(list), defaultdict(list)

    for rangee in ws.iter_rows():
        for cell in rangee:
            valeur = cell.value
            if valeur is None:
                continue
            mois = numero_mois(valeur)
            if mois:
                texte_ligne[cell.row].append((cell.column, mois))
                texte_col[cell.column].append((cell.row, mois))
            elif isinstance(valeur, (datetime, date)):
                date_ligne[cell.row].append((cell.column, valeur.month))
                date_col[cell.column].append((cell.row, valeur.month))

    candidats = []  # (nb_mois, priorite, orientation, axe, positions)
    for source, orientation, priorite in (
        (texte_ligne, "COLONNES", 0), (date_ligne, "COLONNES", 1),
        (texte_col, "LIGNES", 0), (date_col, "LIGNES", 1),
    ):
        for axe, positions in source.items():
            suite = _plus_longue_suite(sorted(positions))
            if len(suite) >= min_mois:
                candidats.append((len(suite), -priorite, orientation == "COLONNES", -axe,
                                  orientation, axe, suite))

    if not candidats:
        return None

    # le plus de mois ; à égalité : texte avant dates, colonnes avant lignes, axe le plus haut/à gauche
    candidats.sort(reverse=True)
    _, _, _, _, orientation, axe, suite = candidats[0]
    positions = {pos: mois for pos, mois in suite}
    derniere = max(positions)

    annuelle = None
    voisin = (ws.cell(row=axe, column=derniere + 1) if orientation == "COLONNES"
              else ws.cell(row=derniere + 1, column=axe))
    if isinstance(voisin.value, str):
        if any(mot in sans_accents(voisin.value).lower() for mot in MOTS_TOTAL):
            annuelle = derniere + 1

    return {"orientation": orientation, "axe_en_tete": axe, "positions": positions,
            "annuelle": annuelle, "origine": "DETECTEE"}


def mise_en_page_forcee(nom_feuille):
    """Mise en page imposée par le fichier de configuration, sinon None."""
    for nom, spec in (charger_config().get("mise_en_page") or {}).items():
        if normaliser_nom(nom) != normaliser_nom(nom_feuille):
            continue
        orientation = spec.get("orientation", "COLONNES").upper()
        debut, fin = spec["mois"].split(":")
        if orientation == "COLONNES":
            a, b = column_index_from_string(debut), column_index_from_string(fin)
            annuelle = column_index_from_string(spec["annuelle"]) if spec.get("annuelle") else None
        else:
            a, b = int(debut), int(fin)
            annuelle = int(spec["annuelle"]) if spec.get("annuelle") else None
        return {"orientation": orientation, "axe_en_tete": int(spec.get("en_tete", 1)),
                "positions": {pos: (i % 12) + 1 for i, pos in enumerate(range(a, b + 1))},
                "annuelle": annuelle, "origine": "CONFIGURATION"}
    return None


def mise_en_page_feuille(ws):
    return mise_en_page_forcee(ws.title) or detecter_mise_en_page(ws)


def mise_en_page_vers_json(mep):
    if mep is None:
        return None
    return {**mep, "positions": {str(k): v for k, v in mep["positions"].items()}}


def mise_en_page_depuis_json(brut):
    if not brut:
        return None
    return {**brut, "positions": {int(k): v for k, v in brut["positions"].items()}}


def charger_mises_en_page():
    """Lit les mises en page enregistrées par analyser_structure ({feuille: mise_en_page})."""
    chemin = dossier_sortie() / "structure.json"
    if not chemin.exists():
        return {}
    with chemin.open("r", encoding="utf-8") as f:
        structure = json.load(f)
    return {nom: mise_en_page_depuis_json(info.get("mise_en_page"))
            for nom, info in structure.get("feuilles", {}).items()}


def nom_fichier_analyse():
    chemin = dossier_sortie() / "structure.json"
    if chemin.exists():
        with chemin.open("r", encoding="utf-8") as f:
            return json.load(f).get("fichier")
    return None


def regroupement_actif(nom_feuille):
    mode = charger_config().get("regroupement_mensuel", "auto")
    if mode == "aucun":
        return False
    if isinstance(mode, list):
        return normaliser_nom(nom_feuille) in {normaliser_nom(n) for n in mode}
    return True


def classer_cellule(mep, colonne, ligne):
    """'MENSUEL', 'ANNUEL' ou None selon la place de la cellule dans la mise en page."""
    if not mep:
        return None
    if mep["orientation"] == "COLONNES":
        if ligne == mep["axe_en_tete"]:
            return None
        pos, annuelle = colonne, mep.get("annuelle")
    else:
        if colonne == mep["axe_en_tete"]:
            return None
        pos, annuelle = ligne, mep.get("annuelle")
    if pos in mep["positions"]:
        return "MENSUEL"
    if annuelle is not None and pos == annuelle:
        return "ANNUEL"
    return None


def axe_de_regroupement(mep, colonne, ligne):
    """Une règle = une ligne (mois en colonnes) ou une colonne (mois en lignes)."""
    return ligne if mep["orientation"] == "COLONNES" else colonne


def decrire_mise_en_page(mep):
    if not mep:
        return "aucune mise en page temporelle détectée"
    cles = sorted(mep["positions"])
    if mep["orientation"] == "COLONNES":
        debut, fin = get_column_letter(cles[0]), get_column_letter(cles[-1])
        texte = f"mois en colonnes {debut}:{fin} (en-tête ligne {mep['axe_en_tete']})"
        if mep.get("annuelle"):
            texte += f", total en {get_column_letter(mep['annuelle'])}"
    else:
        texte = f"mois en lignes {cles[0]}:{cles[-1]} (en-tête colonne {get_column_letter(mep['axe_en_tete'])})"
        if mep.get("annuelle"):
            texte += f", total en ligne {mep['annuelle']}"
    return f"{texte} [{mep.get('origine', 'DETECTEE')}]"


# ============================================================
# DOMAINE / MODELE
# ============================================================

DOMAINES_PAR_DEFAUT = {
    "PRODUCTION": ["PRODUCTION"],
    "EXPORTATION": ["EXPORT"],
    "SALAIRES": ["EMPLOYE", "SALAIRE"],
    "CHARGES": ["CHARGE"],
    "BUDGET": ["SYNTHESE", "BUDGET"],
}


def domaine_de_feuille(nom_feuille):
    """Mots-clés de la configuration, puis mots-clés par défaut, sinon le nom de la feuille."""
    nom = normaliser_nom(nom_feuille)
    if not nom:
        return "AUTRE"
    for table in (charger_config().get("domaines") or {}, DOMAINES_PAR_DEFAUT):
        for domaine, mots in table.items():
            if any(normaliser_nom(m) in nom for m in mots):
                return domaine
    return nom  # ex. "Ventes 2025" -> "VENTES_2025"


def nom_modele():
    config = charger_config()
    if config.get("nom_modele"):
        return config["nom_modele"]
    fichier = nom_fichier_analyse()
    return normaliser_nom(Path(fichier).stem) if fichier else "MODELE_EXCEL"


def version_modele():
    return charger_config().get("version") or f"{nom_modele()}_V1"
