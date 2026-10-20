from pathlib import Path
from openpyxl import load_workbook
import json
import sys

from re_commun import (
    lire_arguments, dossier_sortie, verifier_fichier_excel, numero_mois,
    texte_formule, feuilles_lisibles, mise_en_page_feuille,
    mise_en_page_vers_json, decrire_mise_en_page,
)


def analyser_structure(chemin_fichier):

    fichier = verifier_fichier_excel(chemin_fichier)
    OUTPUT_DIR = dossier_sortie()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print("\n" + "=" * 70)
    print("SXMXDX - REVERSE ENGINEERING")
    print("ANALYSE STRUCTURELLE")
    print("=" * 70)

    wb = load_workbook(fichier, data_only=False, read_only=False)

    resultat = {
        "fichier": fichier.name,
        "nombre_feuilles": len(wb.sheetnames),
        "feuilles": {}
    }

    ignorees = [n for n in wb.sheetnames if n not in {ws.title for ws in feuilles_lisibles(wb)}]
    for nom in ignorees:
        print(f"\n[INFO] Feuille ignorée (graphique, sans cellules) : {nom}")

    for ws in feuilles_lisibles(wb):

        nom_feuille = ws.title
        print(f"\nAnalyse : {nom_feuille}")

        nb_formules = nb_textes = nb_nombres = 0
        mois_detectes = []

        for row in ws.iter_rows():
            for cell in row:
                valeur = cell.value
                if valeur is None:
                    continue

                if texte_formule(valeur):
                    nb_formules += 1
                elif isinstance(valeur, str):
                    nb_textes += 1
                    texte = valeur.strip()
                    if numero_mois(texte) and texte not in mois_detectes:
                        mois_detectes.append(texte)
                elif isinstance(valeur, (int, float)):
                    nb_nombres += 1

        mise_en_page = mise_en_page_feuille(ws)

        resultat["feuilles"][nom_feuille] = {
            "etat": ws.sheet_state,
            "lignes": ws.max_row,
            "colonnes": ws.max_column,
            "formules": nb_formules,
            "textes": nb_textes,
            "nombres": nb_nombres,
            "mois_detectes": mois_detectes,
            "mise_en_page": mise_en_page_vers_json(mise_en_page)
        }

        print(f"  Dimensions : {ws.max_row} lignes x {ws.max_column} colonnes")
        print(f"  Formules   : {nb_formules}")
        print(f"  Mois       : {', '.join(mois_detectes) if mois_detectes else '-'}")
        print(f"  Mise en page : {decrire_mise_en_page(mise_en_page)}")

    wb.close()

    sortie = OUTPUT_DIR / "structure.json"

    with open(sortie, "w", encoding="utf-8") as f:
        json.dump(resultat, f, indent=4, ensure_ascii=False)

    print("\n" + "=" * 70)
    print("[OK] ANALYSE TERMINEE")
    print(f"[OK] Résultat : {sortie}")
    print("=" * 70)


if __name__ == "__main__":

    args = lire_arguments()

    if len(args) != 1:
        print("Utilisation : python analyser_structure.py <fichier_excel> "
              "[--sortie DOSSIER] [--config FICHIER.json]")
        sys.exit(1)

    try:
        analyser_structure(args[0])
    except Exception as e:
        print(f"[ERREUR] {e}")
        sys.exit(1)
