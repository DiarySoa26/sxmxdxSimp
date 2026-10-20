from pathlib import Path
import json
import re
import sys
import unicodedata

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import coordinate_from_string

from re_commun import (
    lire_arguments, dossier_sortie, verifier_fichier_excel, charger_config,
    numero_mois, tous_les_noms_de_mois, sans_accents, mise_en_page_feuille,
    feuilles_lisibles,
)

ARGUMENTS = lire_arguments()

if len(ARGUMENTS) < 1:
    print(
        "Utilisation : "
        "python extraire_semantique.py <fichier_excel> "
        "[--sortie DOSSIER] [--config FICHIER.json]"
    )
    sys.exit(1)

EXCEL_FILE = Path(ARGUMENTS[0]).resolve()

REGLES_FILE = dossier_sortie() / "regles_metier.json"

OUTPUT_FILE = dossier_sortie() / "semantique.json"

# {feuille: mise en page temporelle}, rempli dans main()
MISES_EN_PAGE = {}


# ============================================================
# VALEURS QUI NE SONT PAS DES CONCEPTS METIER
# ============================================================

_VALEURS_DE_BASE = {
    "",
    # devises et unites courantes
    "mga", "usd", "jpy", "eur", "gbp", "chf", "cad", "cny", "zar",
    "ar", "fmg", "xof", "xaf",
    "%", "t", "kg", "g", "tonne", "tonnes", "l", "ml", "m", "m2", "m3",
    "pcs", "pce", "unite", "unites", "unit", "units",
    # agregats generiques
    "total", "total_annuel", "annuel", "mensuel", "annee", "year",
    "monthly", "annual", "yearly", "cumul",
}


def construire_valeurs_non_semantiques():
    valeurs = set(_VALEURS_DE_BASE)
    valeurs |= tous_les_noms_de_mois()
    for extra in charger_config().get("valeurs_non_semantiques", []):
        valeurs.add(normaliser_identifiant(extra) or "")
    return valeurs


# ============================================================
# NORMALISATION
# ============================================================

def normaliser_texte(valeur):

    if valeur is None:
        return None

    texte = str(valeur).strip()

    if not texte:
        return None

    return texte


def normaliser_identifiant(valeur):

    if valeur is None:
        return None

    texte = str(valeur).strip()

    if not texte:
        return None

    # Supprimer les accents
    texte = unicodedata.normalize(
        "NFD",
        texte
    )

    texte = "".join(
        caractere
        for caractere in texte
        if unicodedata.category(caractere) != "Mn"
    )

    texte = texte.lower()

    # Remplacement des caractères non alphanumériques
    texte = re.sub(
        r"[^a-z0-9]+",
        "_",
        texte
    )

    texte = texte.strip("_")

    return texte or None


# ============================================================
# DETECTION DES VALEURS NON SEMANTIQUES
# ============================================================

VALEURS_NON_SEMANTIQUES = construire_valeurs_non_semantiques()

def chercher_libelle_de_ligne(
    ws,
    cellule
):
    """
    Recherche le concept métier d'une ligne.

    Exemple :
        A2 = CA export
        B2 = MGA
        C2:N2 = mois
        O2 = total annuel

    Pour O2, cette fonction retrouve A2 = CA export.
    """

    cellule_obj = ws[cellule]

    ligne = cellule_obj.row

    # Recherche depuis la première colonne
    # jusqu'à la cellule calculée.
    for colonne in range(
        1,
        cellule_obj.column
    ):

        valeur = ws.cell(
            row=ligne,
            column=colonne
        ).value

        texte = normaliser_texte(
            valeur
        )

        if texte is None:
            continue

        if est_libelle_metier(
            texte
        ):

            return {
                "valeur": texte,
                "ligne": ligne,
                "colonne": colonne,
                "position": "DEBUT_LIGNE",
                "distance":
                    cellule_obj.column - colonne
            }

    return None



def chercher_libelle_de_colonne(
    ws,
    cellule
):
    """
    Pendant vertical de chercher_libelle_de_ligne : quand les mois sont
    en lignes, le total annuel d'une colonne herite du libelle en tete de colonne.
    """

    cellule_obj = ws[cellule]

    colonne = cellule_obj.column

    for ligne in range(
        1,
        cellule_obj.row
    ):

        texte = normaliser_texte(
            ws.cell(
                row=ligne,
                column=colonne
            ).value
        )

        if texte is None:
            continue

        if est_libelle_metier(
            texte
        ):

            return {
                "valeur": texte,
                "ligne": ligne,
                "colonne": colonne,
                "position": "DEBUT_COLONNE",
                "distance":
                    cellule_obj.row - ligne
            }

    return None


def est_libelle_metier(valeur):

    texte = normaliser_texte(
        valeur
    )

    if texte is None:
        return False

    # Une formule Excel n'est pas un libellé
    if texte.startswith("="):
        return False

    # Symboles seuls ($, €, %, -...) : unites, pas des concepts
    if not any(c.isalnum() for c in texte):
        return False

    identifiant = normaliser_identifiant(
        texte
    )

    if identifiant is None:
        return False

    if identifiant in VALEURS_NON_SEMANTIQUES:
        return False

    # Mois avec annee : "Janvier 2025", "jan-25"
    if numero_mois(texte):
        return False

    # Nombre seul
    try:
        float(
            texte.replace(",", ".")
        )

        return False

    except ValueError:
        pass

    return True


# ============================================================
# CHARGEMENT DES FICHIERS
# ============================================================

def charger_regles():

    if not REGLES_FILE.exists():

        raise FileNotFoundError(
            f"Fichier de règles introuvable : "
            f"{REGLES_FILE}"
        )

    with REGLES_FILE.open(
        "r",
        encoding="utf-8"
    ) as fichier:

        return json.load(
            fichier
        )


def charger_excel():

    if not EXCEL_FILE.exists():

        raise FileNotFoundError(
            f"Fichier Excel introuvable : "
            f"{EXCEL_FILE}"
        )

    return load_workbook(
        EXCEL_FILE,
        data_only=False,
        read_only=False
    )


# ============================================================
# RECHERCHE HORIZONTALE DU LIBELLE
# ============================================================

def chercher_libelle_gauche(
    ws,
    cellule,
    distance_max=None
):

    if distance_max is None:
        distance_max = charger_config().get(
            "distance_max_libelle",
            10
        )

    cellule_obj = ws[cellule]

    ligne = cellule_obj.row
    colonne = cellule_obj.column

    candidats = []

    for distance in range(
        1,
        distance_max + 1
    ):

        colonne_candidate = (
            colonne - distance
        )

        if colonne_candidate < 1:
            break

        valeur = ws.cell(
            row=ligne,
            column=colonne_candidate
        ).value

        texte = normaliser_texte(
            valeur
        )

        if texte is None:
            continue

        candidats.append({
            "valeur": texte,
            "ligne": ligne,
            "colonne": colonne_candidate,
            "position": "GAUCHE",
            "distance": distance,
            "est_semantique":
                est_libelle_metier(texte)
        })

        # IMPORTANT :
        # on ne s'arrête plus sur MGA, %, Total annuel, etc.
        if est_libelle_metier(
            texte
        ):

            return {
                "valeur": texte,
                "ligne": ligne,
                "colonne": colonne_candidate,
                "position": "GAUCHE",
                "distance": distance,
                "candidats": candidats
            }

    return None


# ============================================================
# RECHERCHE VERTICALE
# ============================================================

def chercher_libelle_haut(
    ws,
    cellule,
    distance_max=None
):

    if distance_max is None:
        distance_max = charger_config().get(
            "distance_max_libelle",
            10
        )

    cellule_obj = ws[cellule]

    ligne = cellule_obj.row
    colonne = cellule_obj.column

    candidats = []

    for distance in range(
        1,
        distance_max + 1
    ):

        ligne_candidate = (
            ligne - distance
        )

        if ligne_candidate < 1:
            break

        valeur = ws.cell(
            row=ligne_candidate,
            column=colonne
        ).value

        texte = normaliser_texte(
            valeur
        )

        if texte is None:
            continue

        candidats.append({
            "valeur": texte,
            "ligne": ligne_candidate,
            "colonne": colonne,
            "position": "HAUT",
            "distance": distance,
            "est_semantique":
                est_libelle_metier(texte)
        })

        if est_libelle_metier(
            texte
        ):

            return {
                "valeur": texte,
                "ligne": ligne_candidate,
                "colonne": colonne,
                "position": "HAUT",
                "distance": distance,
                "candidats": candidats
            }

    return None


# ============================================================
# RECHERCHE GENERALE
# ============================================================

def chercher_libelle(
    ws,
    cellule
):

    # --------------------------------------------------------
    # PRIORITE 1 :
    # chercher sur la même ligne à gauche
    #
    # C'est généralement là que se trouve :
    #
    # Chiffre d'affaires | MGA | Janvier ...
    # Charges            | MGA | Janvier ...
    # Taux de marge      | %   | Janvier ...
    # --------------------------------------------------------

    gauche = chercher_libelle_gauche(
        ws,
        cellule
    )

    if gauche is not None:
        return gauche

    # --------------------------------------------------------
    # PRIORITE 2 :
    # chercher au-dessus
    # --------------------------------------------------------

    haut = chercher_libelle_haut(
        ws,
        cellule
    )

    if haut is not None:
        return haut

    return None


def chercher_libelle_selon_mise_en_page(
    ws,
    cellule,
    mise_en_page
):
    """
    Cellule situee dans la zone des mois (ou dans la colonne/ligne de total) :
    le libelle est celui de la ligne (mois en colonnes) ou de la colonne
    (mois en lignes), QUELLE QUE SOIT la distance qui les separe.
    Retourne None si la cellule n'est pas dans cette zone.
    """

    if not mise_en_page:
        return None

    from re_commun import classer_cellule
    from openpyxl.utils import column_index_from_string

    colonne, ligne = coordinate_from_string(
        cellule
    )

    categorie = classer_cellule(
        mise_en_page,
        column_index_from_string(colonne.upper()),
        int(ligne)
    )

    if categorie is None:
        return None

    cellule_obj = ws[cellule]

    colonnes = mise_en_page["orientation"] == "COLONNES"

    if categorie == "ANNUEL":

        # Meme concept metier que les mois de la ligne / colonne
        if colonnes:
            return chercher_libelle_de_ligne(ws, cellule)

        return chercher_libelle_de_colonne(ws, cellule)

    # MENSUEL
    if colonnes:
        return chercher_libelle_gauche(
            ws,
            cellule,
            distance_max=cellule_obj.column
        )

    return chercher_libelle_haut(
        ws,
        cellule,
        distance_max=cellule_obj.row
    )


# ============================================================
# PERIODICITE SEMANTIQUE
# ============================================================

def determiner_periodicite_semantique(
    feuille,
    cellule,
    periodicite_technique
):

    if not feuille or not cellule:
        return periodicite_technique

    mise_en_page = MISES_EN_PAGE.get(feuille)

    if not mise_en_page:
        return periodicite_technique

    colonne, ligne = coordinate_from_string(
        cellule
    )

    from re_commun import classer_cellule
    from openpyxl.utils import column_index_from_string

    categorie = classer_cellule(
        mise_en_page,
        column_index_from_string(colonne.upper()),
        int(ligne)
    )

    # Mois detectes dans les en-tetes -> MENSUELLE
    # Colonne (ou ligne) de total juste apres les mois -> ANNUELLE
    if categorie == "MENSUEL":
        return "MENSUELLE"

    if categorie == "ANNUEL":
        return "ANNUELLE"

    return periodicite_technique


# ============================================================
# TYPE SEMANTIQUE
# ============================================================

def determiner_type_semantique(
    periodicite
):

    if periodicite == "MENSUELLE":
        return "MESURE_MENSUELLE"

    if periodicite == "ANNUELLE":
        return "MESURE_ANNUELLE"

    return "MESURE"


# ============================================================
# DEPENDANCES SEMANTIQUES ET EXPRESSION METIER
#
# Remplace le mapping manuel (RG_xxx -> sources/expression) :
# chaque cellule source est traduite en concept metier grace a
# son libelle dans le classeur.
# ============================================================

import re as _re

_CELLULE_SIMPLE = _re.compile(r"^[A-Za-z]{1,3}\d+$")

_CACHE_LIBELLES = {}


def libelle_de_source(workbook, feuille, plage):

    cle = (feuille, plage)

    if cle in _CACHE_LIBELLES:
        return _CACHE_LIBELLES[cle]

    resultat = None

    try:

        if feuille in workbook.sheetnames:

            ws = workbook[feuille]

            if hasattr(ws, "iter_rows"):

                premiere_cellule = plage.split(":")[0]

                if _CELLULE_SIMPLE.match(premiere_cellule):

                    resultat = chercher_libelle_selon_mise_en_page(
                        ws,
                        premiere_cellule,
                        MISES_EN_PAGE.get(feuille)
                    )

                if resultat is not None:

                    pass

                elif ":" in plage:

                    debut, fin = plage.split(":")

                    if (
                        _CELLULE_SIMPLE.match(debut)
                        and _CELLULE_SIMPLE.match(fin)
                    ):

                        premiere = ws[debut]
                        derniere = ws[fin]

                        # Colonne de donnees : l'en-tete est au-dessus
                        if (
                            premiere.column == derniere.column
                            and premiere.row != derniere.row
                        ):

                            resultat = (
                                chercher_libelle_haut(ws, debut)
                                or chercher_libelle_gauche(ws, debut)
                            )

                        else:

                            resultat = chercher_libelle(ws, debut)

                elif _CELLULE_SIMPLE.match(plage):

                    resultat = chercher_libelle(ws, plage)

    except Exception:

        resultat = None

    _CACHE_LIBELLES[cle] = resultat

    return resultat


def construire_expression(
    workbook,
    formule,
    feuille_courante
):
    """
    '=C2-C3-C4' -> 'chiffre_affaires - charges - frais'
    Retourne None si une des sources n'a pas de libelle.
    """

    if not formule:
        return None

    try:
        tokens = Tokenizer(formule).items
    except Exception:
        return None

    morceaux = []

    for token in tokens:

        if (
            token.type == "OPERAND"
            and token.subtype == "RANGE"
        ):

            valeur = token.value.replace("$", "").replace("'", "")

            if "!" in valeur:
                feuille_source, cellule = valeur.rsplit("!", 1)
            else:
                feuille_source, cellule = feuille_courante, valeur

            libelle = libelle_de_source(
                workbook,
                feuille_source,
                cellule
            )

            if not libelle:
                return None

            identifiant = normaliser_identifiant(
                libelle["valeur"]
            )

            if feuille_source != feuille_courante:
                identifiant = (
                    f"{normaliser_identifiant(feuille_source)}."
                    f"{identifiant}"
                )

            morceaux.append(identifiant)

        elif token.type == "OPERATOR-INFIX":
            morceaux.append(f" {token.value} ")

        elif token.type == "SEP":
            morceaux.append(", ")

        elif token.type == "WHITE-SPACE":
            continue

        else:
            morceaux.append(token.value)

    return "".join(morceaux).strip()


def analyser_dependances_semantiques(
    workbook,
    regle,
    feuille_courante
):

    dependances = []

    for dep in regle.get("dependances_reference") or []:

        libelle = libelle_de_source(
            workbook,
            dep.get("feuille"),
            dep.get("cellule")
        )

        dependances.append({
            "feuille": dep.get("feuille"),
            "cellule": dep.get("cellule"),
            "libelle": (
                libelle["valeur"] if libelle else None
            ),
            "identifiant_semantique": (
                normaliser_identifiant(libelle["valeur"])
                if libelle else None
            )
        })

    expression = construire_expression(
        workbook,
        regle.get("formule_excel_reference"),
        feuille_courante
    )

    return dependances, expression


# ============================================================
# ANALYSE D'UNE REGLE
# ============================================================

def analyser_regle(
    workbook,
    regle
):

    feuille = regle.get(
        "feuille"
    )

    cellule = regle.get(
        "cellule_cible_reference"
    )

    periodicite_technique = (
        regle.get("periodicite")
    )

    periodicite_semantique = (
        determiner_periodicite_semantique(
            feuille,
            cellule,
            periodicite_technique
        )
    )

    resultat = {

        "regle_technique":
            regle.get("id"),

        "domaine":
            regle.get("domaine"),

        "feuille":
            feuille,

        "cellule_reference":
            cellule,

        "operation":
            regle.get("operation"),

        "periodicite_technique":
            periodicite_technique,

        "periodicite":
            periodicite_semantique,

        "nombre_occurrences":
            regle.get(
                "nombre_occurrences",
                1
            ),

        "formule":
            regle.get(
                "formule_excel_reference"
            ),

        "dependances":
            regle.get(
                "dependances_reference",
                []
            )
    }

    # --------------------------------------------------------
    # FEUILLE ABSENTE
    # --------------------------------------------------------

    if feuille not in workbook.sheetnames:

        resultat["libelle_excel"] = None

        resultat[
            "identifiant_semantique"
        ] = None

        resultat["statut"] = (
            "FEUILLE_INTROUVABLE"
        )

        return resultat

    ws = workbook[
        feuille
    ]

    (
        resultat["dependances_semantiques"],
        resultat["expression_metier"]
    ) = analyser_dependances_semantiques(
        workbook,
        regle,
        feuille
    )

    # --------------------------------------------------------
    # RECHERCHE DU LIBELLE
    # --------------------------------------------------------

    # ========================================================
    # RECHERCHE SEMANTIQUE
    # ========================================================

    # Cellule dans la zone des mois / du total : libelle de la ligne
    # (ou de la colonne) sans limite de distance.
    # Sinon : recherche a gauche puis au-dessus.
    libelle = chercher_libelle_selon_mise_en_page(
        ws,
        cellule,
        MISES_EN_PAGE.get(feuille)
    )

    if libelle is None:

        libelle = chercher_libelle(
            ws,
            cellule
        )

    if libelle is None:

        resultat["libelle_excel"] = None

        resultat[
            "identifiant_semantique"
        ] = None

        resultat["statut"] = (
            "A_VALIDER"
        )

        return resultat

    # --------------------------------------------------------
    # LIBELLE TROUVE
    # --------------------------------------------------------

    texte = libelle[
        "valeur"
    ]

    resultat["libelle_excel"] = (
        texte
    )

    resultat[
        "identifiant_semantique"
    ] = normaliser_identifiant(
        texte
    )

    resultat["position_libelle"] = (
        libelle["position"]
    )

    resultat["distance_libelle"] = (
        libelle["distance"]
    )

    resultat["type_semantique"] = (
        determiner_type_semantique(
            periodicite_semantique
        )
    )

    resultat["statut"] = (
        "SEMANTIQUE_DETECTEE"
    )

    return resultat


# ============================================================
# VALIDATION DES DOUBLONS SEMANTIQUES
# ============================================================

def analyser_doublons(
    resultats
):

    groupes = {}

    for regle in resultats:

        cible = regle.get(
            "identifiant_semantique"
        )

        periodicite = regle.get(
            "periodicite"
        )

        if not cible:
            continue

        cle = (
            cible,
            periodicite
        )

        groupes.setdefault(
            cle,
            []
        ).append(
            regle["regle_technique"]
        )

    doublons = []

    for cle, ids in groupes.items():

        if len(ids) > 1:

            doublons.append({
                "identifiant_semantique":
                    cle[0],

                "periodicite":
                    cle[1],

                "regles":
                    ids
            })

    return doublons


# ============================================================
# STATISTIQUES
# ============================================================

def calculer_statistiques(
    resultats
):

    total = len(
        resultats
    )

    detectees = sum(
        1
        for regle in resultats
        if regle.get("statut")
        == "SEMANTIQUE_DETECTEE"
    )

    a_valider = (
        total - detectees
    )

    mensuelles = sum(
        1
        for regle in resultats
        if regle.get("periodicite")
        == "MENSUELLE"
    )

    annuelles = sum(
        1
        for regle in resultats
        if regle.get("periodicite")
        == "ANNUELLE"
    )

    uniques = sum(
        1
        for regle in resultats
        if regle.get("periodicite")
        == "UNIQUE"
    )

    return {
        "nombre_regles":
            total,

        "semantique_detectee":
            detectees,

        "a_valider":
            a_valider,

        "periodicites": {
            "MENSUELLE":
                mensuelles,

            "ANNUELLE":
                annuelles,

            "UNIQUE":
                uniques,

            "REPETEE":
                total - mensuelles - annuelles - uniques
        }
    }


# ============================================================
# PROGRAMME PRINCIPAL
# ============================================================

def main():

    print()
    print("=" * 75)
    print(
        "SXMXDX - EXTRACTION SEMANTIQUE DES REGLES"
    )
    print("=" * 75)

    # --------------------------------------------------------
    # CHARGEMENT
    # --------------------------------------------------------

    document = charger_regles()

    workbook = charger_excel()

    global MISES_EN_PAGE

    MISES_EN_PAGE = {
        ws.title: mise_en_page_feuille(ws)
        for ws in feuilles_lisibles(workbook)
    }

    MISES_EN_PAGE = {
        nom: mep
        for nom, mep in MISES_EN_PAGE.items()
        if mep
    }

    regles = document.get(
        "regles",
        []
    )

    print(
        f"\nRègles techniques à analyser : "
        f"{len(regles)}"
    )

    print()

    # --------------------------------------------------------
    # ANALYSE
    # --------------------------------------------------------

    resultats = []

    for regle in regles:

        resultat = analyser_regle(
            workbook,
            regle
        )

        resultats.append(
            resultat
        )

        id_regle = (
            resultat.get(
                "regle_technique"
            )
        )

        libelle = (
            resultat.get(
                "libelle_excel"
            )
            or "NON DETECTE"
        )

        periodicite = (
            resultat.get(
                "periodicite"
            )
        )

        statut = (
            resultat.get(
                "statut"
            )
        )

        print(
            f"{id_regle:<25} "
            f"{libelle:<35} "
            f"{periodicite:<12} "
            f"{statut}"
        )

    # --------------------------------------------------------
    # STATISTIQUES
    # --------------------------------------------------------

    statistiques = (
        calculer_statistiques(
            resultats
        )
    )

    doublons = (
        analyser_doublons(
            resultats
        )
    )

    # --------------------------------------------------------
    # DOCUMENT FINAL
    # --------------------------------------------------------

    sortie = {

        "modele":
            document.get(
                "modele"
            ),

        "version":
            document.get(
                "version"
            ),

        "origine":
            "EXTRACTION_SEMANTIQUE_EXCEL",

        "source_excel":
            str(EXCEL_FILE),

        "source_regles":
            str(REGLES_FILE),

        "statistiques":
            statistiques,

        "doublons_semantiques":
            doublons,

        "regles":
            resultats
    }

    # --------------------------------------------------------
    # ECRITURE
    # --------------------------------------------------------

    OUTPUT_FILE.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    with OUTPUT_FILE.open(
        "w",
        encoding="utf-8"
    ) as fichier:

        json.dump(
            sortie,
            fichier,
            ensure_ascii=False,
            indent=4
        )

    # --------------------------------------------------------
    # RESULTAT
    # --------------------------------------------------------

    print()
    print("-" * 75)
    print("RESULTAT")
    print("-" * 75)

    print(
        f"Règles analysées      : "
        f"{statistiques['nombre_regles']}"
    )

    print(
        f"Sémantiques détectées : "
        f"{statistiques['semantique_detectee']}"
    )

    print(
        f"À valider             : "
        f"{statistiques['a_valider']}"
    )

    print()
    print("Périodicités :")

    print(
        f"  MENSUELLE : "
        f"{statistiques['periodicites']['MENSUELLE']}"
    )

    print(
        f"  ANNUELLE  : "
        f"{statistiques['periodicites']['ANNUELLE']}"
    )

    print(
        f"  UNIQUE    : "
        f"{statistiques['periodicites']['UNIQUE']}"
    )

    if doublons:

        print()
        print(
            "[INFO] Concepts identiques détectés "
            "dans plusieurs règles :"
        )

        for doublon in doublons:

            print(
                f"  {doublon['identifiant_semantique']} "
                f"({doublon['periodicite']}) "
                f"→ {', '.join(doublon['regles'])}"
            )

    print()
    print(
        f"[OK] Fichier généré : "
        f"{OUTPUT_FILE}"
    )

    print("=" * 75)


if __name__ == "__main__":
    main()