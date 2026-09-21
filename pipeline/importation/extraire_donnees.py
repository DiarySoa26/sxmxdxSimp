from pathlib import Path
from openpyxl import load_workbook
import csv
import re
import sys


OUTPUT_BASE = Path("data/staging")


def nom_fichier_propre(nom):

    nom = nom.strip().lower()

    nom = (
        nom.replace("é", "e")
           .replace("è", "e")
           .replace("ê", "e")
           .replace("à", "a")
           .replace("ù", "u")
    )

    nom = re.sub(r"[^a-z0-9]+", "_", nom)

    return nom.strip("_")


def extraire_donnees(chemin_fichier):

    fichier = Path(chemin_fichier)

    if not fichier.exists():
        raise FileNotFoundError(fichier)

    dossier_import = OUTPUT_BASE / fichier.stem

    dossier_import.mkdir(
        parents=True,
        exist_ok=True
    )

    print("\n" + "=" * 70)
    print("SXMXDX - EXTRACTION VERS STAGING")
    print("=" * 70)

    # data_only=True :
    # on récupère les valeurs calculées visibles dans Excel
    wb = load_workbook(
        fichier,
        data_only=True
    )

    for nom_feuille in wb.sheetnames:

        ws = wb[nom_feuille]

        nom_csv = (
            nom_fichier_propre(nom_feuille)
            + ".csv"
        )

        sortie = dossier_import / nom_csv

        with open(
            sortie,
            "w",
            newline="",
            encoding="utf-8"
        ) as f:

            writer = csv.writer(f)

            for row in ws.iter_rows(
                values_only=True
            ):

                writer.writerow(row)

        print(
            f"[OK] {nom_feuille:<25} "
            f"-> {sortie}"
        )

    wb.close()

    print("\n" + "-" * 70)
    print("EXTRACTION TERMINEE")
    print(f"STAGING : {dossier_import}")

    return dossier_import


if __name__ == "__main__":

    if len(sys.argv) != 2:

        print(
            "Utilisation : python extraire_donnees.py "
            "<fichier_excel>"
        )

        sys.exit(1)

    try:
        extraire_donnees(sys.argv[1])

    except Exception as e:
        print(f"[ERREUR] {e}")
        sys.exit(1)