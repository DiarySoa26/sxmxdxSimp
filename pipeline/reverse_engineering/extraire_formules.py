from pathlib import Path
from openpyxl import load_workbook
import csv
import sys

from re_commun import (
    lire_arguments, dossier_sortie, verifier_fichier_excel,
    texte_formule, feuilles_lisibles,
)


def extraire_formules(chemin_fichier):

    fichier = verifier_fichier_excel(chemin_fichier)
    OUTPUT_DIR = dossier_sortie()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print("\n" + "=" * 70)
    print("SXMXDX - EXTRACTION DES FORMULES")
    print("=" * 70)

    wb = load_workbook(fichier, data_only=False)

    formules = []

    for ws in feuilles_lisibles(wb):

        compteur = 0

        for row in ws.iter_rows():
            for cell in row:
                formule = texte_formule(cell.value)
                if formule:
                    formules.append({
                        "feuille": ws.title,
                        "cellule": cell.coordinate,
                        "formule": formule
                    })
                    compteur += 1

        print(f"{ws.title:<25} : {compteur} formule(s)")

    wb.close()

    sortie = OUTPUT_DIR / "formules.csv"

    with open(sortie, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["feuille", "cellule", "formule"])
        writer.writeheader()
        writer.writerows(formules)

    print("\n" + "-" * 70)
    print(f"Total : {len(formules)} formule(s)")
    if not formules:
        print("[ATTENTION] Aucune formule trouvée : ce classeur ne contient que des valeurs.")
    print(f"[OK] Fichier généré : {sortie}")


if __name__ == "__main__":

    args = lire_arguments()

    if len(args) != 1:
        print("Utilisation : python extraire_formules.py <fichier_excel> "
              "[--sortie DOSSIER] [--config FICHIER.json]")
        sys.exit(1)

    try:
        extraire_formules(args[0])
    except Exception as e:
        print(f"[ERREUR] {e}")
        sys.exit(1)
