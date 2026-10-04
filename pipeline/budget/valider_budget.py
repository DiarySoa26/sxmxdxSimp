# from decimal import Decimal, InvalidOperation
# from pathlib import Path
# import sys

# import psycopg2
# from psycopg2.extras import RealDictCursor
# from openpyxl import load_workbook


# # ============================================================
# # PROJET
# # ============================================================

# PROJECT_ROOT = Path(__file__).resolve().parents[2]

# if str(PROJECT_ROOT) not in sys.path:
#     sys.path.insert(0, str(PROJECT_ROOT))


# # ============================================================
# # CONFIGURATION POSTGRESQL
# # ============================================================

# DB_CONFIG = {
#     "host": "localhost",
#     "port": 5434,
#     "dbname": "sxmxdx2",
#     "user": "sxmxdx",
#     "password": "diary",
# }


# # ============================================================
# # CONFIGURATION VALIDATION
# # ============================================================

# FEUILLE_SYNTHESE = "Synthese"

# # Tolérance monétaire en MGA
# TOLERANCE_MONTANT = Decimal("1.00")

# # Tolérance marge :
# # 0.0001 = 0.01 point de pourcentage
# TOLERANCE_MARGE = Decimal("0.0001")


# # ============================================================
# # CONNEXION
# # ============================================================

# def connecter():

#     return psycopg2.connect(
#         **DB_CONFIG
#     )


# # ============================================================
# # CONVERSION DECIMAL
# # ============================================================

# def decimaliser(valeur):

#     if valeur is None:
#         return Decimal("0")

#     if isinstance(valeur, Decimal):
#         return valeur

#     try:
#         return Decimal(str(valeur))

#     except (InvalidOperation, ValueError, TypeError):

#         raise ValueError(
#             f"Valeur numérique invalide : {valeur}"
#         )


# # ============================================================
# # DERNIER BUDGET GENERE
# # ============================================================

# def recuperer_dernier_budget(connexion):

#     with connexion.cursor(
#         cursor_factory=RealDictCursor
#     ) as curseur:

#         curseur.execute(
#             """
#             SELECT
#                 b.id AS budget_id,
#                 b.import_id,
#                 b.exercice_id,
#                 b.ca_export,
#                 b.charges_exploitation,
#                 b.frais_export,
#                 b.masse_salariale,
#                 b.resultat_operationnel,
#                 b.marge_operationnelle,
#                 b.statut,
#                 e.annee
#             FROM budget b

#             JOIN exercice e
#                 ON e.id = b.exercice_id

#             WHERE b.statut = 'GENERE'
#               AND b.type_budget = 'GENERE'

#             ORDER BY b.id DESC

#             LIMIT 1
#             """
#         )

#         budget = curseur.fetchone()

#     if budget is None:

#         raise ValueError(
#             "Aucun budget GENERE trouvé."
#         )

#     return budget


# # ============================================================
# # IMPORT
# # ============================================================

# def recuperer_import(
#     connexion,
#     import_id
# ):

#     with connexion.cursor(
#         cursor_factory=RealDictCursor
#     ) as curseur:

#         curseur.execute(
#             """
#             SELECT *
#             FROM import_fichier
#             WHERE id = %s
#             """,
#             (import_id,)
#         )

#         ligne = curseur.fetchone()

#     if ligne is None:

#         raise ValueError(
#             f"Import {import_id} introuvable."
#         )

#     return ligne


# # ============================================================
# # BUDGETS MENSUELS
# # ============================================================

# def recuperer_budget_mensuel(
#     connexion,
#     budget_id
# ):

#     with connexion.cursor(
#         cursor_factory=RealDictCursor
#     ) as curseur:

#         curseur.execute(
#             """
#             SELECT
#                 bm.*,
#                 m.nom AS mois_nom
#             FROM budget_mensuel bm

#             JOIN mois m
#                 ON m.id = bm.mois_id

#             WHERE bm.budget_id = %s

#             ORDER BY bm.mois_id
#             """,
#             (budget_id,)
#         )

#         lignes = curseur.fetchall()

#     if len(lignes) != 12:

#         raise ValueError(
#             "Le budget doit contenir exactement "
#             f"12 mois. Trouvé : {len(lignes)}"
#         )

#     return lignes


# # ============================================================
# # RECHERCHE DU FICHIER EXCEL
# # ============================================================

# def trouver_fichier_excel(import_fichier):

#     """
#     Cherche le chemin du fichier dans les colonnes
#     disponibles de import_fichier.

#     Comme le nom exact de la colonne peut dépendre
#     de la version du schéma, plusieurs noms sont testés.
#     """

#     colonnes_possibles = [
#         "chemin_fichier",
#         "fichier",
#         "nom_fichier",
#         "filepath",
#         "file_path",
#     ]

#     valeur = None

#     for colonne in colonnes_possibles:

#         if colonne in import_fichier:

#             if import_fichier[colonne]:

#                 valeur = str(
#                     import_fichier[colonne]
#                 )

#                 break

#     if valeur is None:

#         raise ValueError(
#             "Impossible de retrouver le fichier Excel "
#             "dans import_fichier. "
#             "Aucune colonne de chemin reconnue."
#         )

#     chemin = Path(valeur)

#     # Chemin absolu valide
#     if chemin.is_absolute() and chemin.exists():

#         return chemin

#     # Chemin relatif au projet
#     candidat = (
#         PROJECT_ROOT /
#         chemin
#     ).resolve()

#     if candidat.exists():

#         return candidat

#     # Seulement le nom du fichier :
#     # recherche dans data/uploads
#     upload_dir = (
#         PROJECT_ROOT /
#         "data" /
#         "uploads"
#     )

#     if upload_dir.exists():

#         correspondances = list(
#             upload_dir.rglob(
#                 chemin.name
#             )
#         )

#         if correspondances:

#             correspondances.sort(
#                 key=lambda p: p.stat().st_mtime,
#                 reverse=True
#             )

#             return correspondances[0]

#     raise FileNotFoundError(
#         "Fichier Excel source introuvable : "
#         f"{valeur}"
#     )


# # ============================================================
# # NORMALISATION TEXTE
# # ============================================================

# def normaliser_texte(valeur):

#     if valeur is None:
#         return ""

#     return (
#         str(valeur)
#         .strip()
#         .lower()
#         .replace("é", "e")
#         .replace("è", "e")
#         .replace("ê", "e")
#         .replace("à", "a")
#         .replace("ô", "o")
#         .replace("û", "u")
#         .replace("ï", "i")
#     )


# # ============================================================
# # RECHERCHE D'UNE LIGNE DANS SYNTHESE
# # ============================================================

# def trouver_ligne_indicateur(
#     feuille,
#     mots_cles
# ):

#     mots_cles = [
#         normaliser_texte(mot)
#         for mot in mots_cles
#     ]

#     for ligne in feuille.iter_rows():

#         for cellule in ligne:

#             texte = normaliser_texte(
#                 cellule.value
#             )

#             if not texte:
#                 continue

#             if all(
#                 mot in texte
#                 for mot in mots_cles
#             ):

#                 return cellule.row

#     return None


# # ============================================================
# # EXTRACTION D'UNE LIGNE DE SYNTHESE
# # ============================================================

# def extraire_ligne_synthese(
#     feuille,
#     ligne
# ):

#     """
#     Modèle actuel de Synthese :
#     C:N = Janvier à Décembre
#     O   = Total annuel
#     """

#     mensuel = []

#     for colonne in range(3, 15):

#         valeur = feuille.cell(
#             row=ligne,
#             column=colonne
#         ).value

#         mensuel.append(
#             decimaliser(valeur)
#         )

#     annuel = decimaliser(
#         feuille.cell(
#             row=ligne,
#             column=15
#         ).value
#     )

#     return {
#         "mensuel": mensuel,
#         "annuel": annuel,
#     }


# # ============================================================
# # EXTRACTION SYNTHESE EXCEL
# # ============================================================

# def extraire_synthese_excel(
#     fichier_excel
# ):

#     classeur = load_workbook(
#         fichier_excel,
#         data_only=True,
#         read_only=True
#     )

#     try:

#         if FEUILLE_SYNTHESE not in classeur.sheetnames:

#             raise ValueError(
#                 f"Feuille '{FEUILLE_SYNTHESE}' "
#                 "introuvable."
#             )

#         feuille = classeur[
#             FEUILLE_SYNTHESE
#         ]

#         definitions = {
#             "ca_export": [
#                 "ca",
#                 "export"
#             ],

#             "charges_exploitation": [
#                 "charges",
#                 "exploitation"
#             ],

#             "frais_export": [
#                 "frais",
#                 "export"
#             ],

#             "masse_salariale": [
#                 "masse",
#                 "salariale"
#             ],

#             "resultat_operationnel": [
#                 "resultat",
#                 "operationnel"
#             ],

#             "marge_operationnelle": [
#                 "marge",
#                 "operationnelle"
#             ],
#         }

#         resultat = {}

#         for indicateur, mots_cles in definitions.items():

#             ligne = trouver_ligne_indicateur(
#                 feuille,
#                 mots_cles
#             )

#             if ligne is None:

#                 raise ValueError(
#                     "Indicateur introuvable "
#                     "dans Synthese : "
#                     f"{indicateur}"
#                 )

#             resultat[indicateur] = (
#                 extraire_ligne_synthese(
#                     feuille,
#                     ligne
#                 )
#             )

#         return resultat

#     finally:

#         classeur.close()


# # ============================================================
# # COMPARAISON
# # ============================================================

# def comparer(
#     attendu,
#     obtenu,
#     tolerance
# ):

#     attendu = decimaliser(
#         attendu
#     )

#     obtenu = decimaliser(
#         obtenu
#     )

#     ecart = abs(
#         attendu - obtenu
#     )

#     valide = (
#         ecart <= tolerance
#     )

#     return {
#         "attendu": attendu,
#         "obtenu": obtenu,
#         "ecart": ecart,
#         "valide": valide,
#     }


# # ============================================================
# # AFFICHAGE MONTANT
# # ============================================================

# def afficher_nombre(valeur):

#     valeur = decimaliser(
#         valeur
#     )

#     return f"{valeur:,.2f}"


# # ============================================================
# # VALIDATION
# # ============================================================

# def valider_budget():

#     connexion = None

#     try:

#         print()
#         print("=" * 75)
#         print(
#             "SXMXDX - VALIDATION DU BUDGET"
#         )
#         print("=" * 75)

#         # ----------------------------------------------------
#         # 1. CONNEXION
#         # ----------------------------------------------------

#         connexion = connecter()

#         # ----------------------------------------------------
#         # 2. DERNIER BUDGET
#         # ----------------------------------------------------

#         budget = recuperer_dernier_budget(
#             connexion
#         )

#         budget_id = budget[
#             "budget_id"
#         ]

#         import_id = budget[
#             "import_id"
#         ]

#         exercice_id = budget[
#             "exercice_id"
#         ]

#         annee = budget[
#             "annee"
#         ]

#         print()
#         print(
#             f"Budget ID   : {budget_id}"
#         )

#         print(
#             f"Import ID   : {import_id}"
#         )

#         print(
#             f"Exercice ID : {exercice_id}"
#         )

#         print(
#             f"Année       : {annee}"
#         )

#         # ----------------------------------------------------
#         # 3. IMPORT
#         # ----------------------------------------------------

#         import_fichier = recuperer_import(
#             connexion,
#             import_id
#         )

#         # ----------------------------------------------------
#         # 4. FICHIER EXCEL
#         # ----------------------------------------------------

#         fichier_excel = trouver_fichier_excel(
#             import_fichier
#         )

#         print(
#             f"Excel       : {fichier_excel}"
#         )

#         # ----------------------------------------------------
#         # 5. BUDGET MENSUEL
#         # ----------------------------------------------------

#         budgets_mensuels = (
#             recuperer_budget_mensuel(
#                 connexion,
#                 budget_id
#             )
#         )

#         # ----------------------------------------------------
#         # 6. EXCEL
#         # ----------------------------------------------------

#         excel = extraire_synthese_excel(
#             fichier_excel
#         )

#         # ----------------------------------------------------
#         # 7. INDICATEURS
#         # ----------------------------------------------------

#         indicateurs = [
#             "ca_export",
#             "charges_exploitation",
#             "frais_export",
#             "masse_salariale",
#             "resultat_operationnel",
#             "marge_operationnelle",
#         ]

#         erreurs = []

#         # ----------------------------------------------------
#         # 8. VALIDATION MENSUELLE
#         # ----------------------------------------------------

#         print()
#         print("=" * 75)
#         print(
#             "VALIDATION MENSUELLE"
#         )
#         print("=" * 75)

#         for index, mois in enumerate(
#             budgets_mensuels
#         ):

#             print()
#             print(
#                 f"[{mois['mois_nom']}]"
#             )

#             for indicateur in indicateurs:

#                 attendu = (
#                     excel[
#                         indicateur
#                     ][
#                         "mensuel"
#                     ][
#                         index
#                     ]
#                 )

#                 obtenu = mois[
#                     indicateur
#                 ]

#                 if indicateur == "marge_operationnelle":

#                     tolerance = (
#                         TOLERANCE_MARGE
#                     )

#                 else:

#                     tolerance = (
#                         TOLERANCE_MONTANT
#                     )

#                 comparaison = comparer(
#                     attendu,
#                     obtenu,
#                     tolerance
#                 )

#                 statut = (
#                     "OK"
#                     if comparaison[
#                         "valide"
#                     ]
#                     else "ECART"
#                 )

#                 print(
#                     f"  {indicateur:<25} "
#                     f"{statut:<5} "
#                     f"Excel={afficher_nombre(comparaison['attendu'])} "
#                     f"DB={afficher_nombre(comparaison['obtenu'])} "
#                     f"Ecart={afficher_nombre(comparaison['ecart'])}"
#                 )

#                 if not comparaison[
#                     "valide"
#                 ]:

#                     erreurs.append(
#                         {
#                             "periode":
#                                 mois[
#                                     "mois_nom"
#                                 ],

#                             "indicateur":
#                                 indicateur,

#                             **comparaison,
#                         }
#                     )

#         # ----------------------------------------------------
#         # 9. VALIDATION ANNUELLE
#         # ----------------------------------------------------

#         print()
#         print("=" * 75)
#         print(
#             "VALIDATION ANNUELLE"
#         )
#         print("=" * 75)

#         for indicateur in indicateurs:

#             attendu = excel[
#                 indicateur
#             ][
#                 "annuel"
#             ]

#             obtenu = budget[
#                 indicateur
#             ]

#             if indicateur == "marge_operationnelle":

#                 tolerance = (
#                     TOLERANCE_MARGE
#                 )

#             else:

#                 tolerance = (
#                     TOLERANCE_MONTANT
#                 )

#             comparaison = comparer(
#                 attendu,
#                 obtenu,
#                 tolerance
#             )

#             statut = (
#                 "OK"
#                 if comparaison[
#                     "valide"
#                 ]
#                 else "ECART"
#             )

#             print(
#                 f"{indicateur:<25} "
#                 f"{statut:<5} "
#                 f"Excel={afficher_nombre(comparaison['attendu'])} "
#                 f"DB={afficher_nombre(comparaison['obtenu'])} "
#                 f"Ecart={afficher_nombre(comparaison['ecart'])}"
#             )

#             if not comparaison[
#                 "valide"
#             ]:

#                 erreurs.append(
#                     {
#                         "periode":
#                             "ANNUEL",

#                         "indicateur":
#                             indicateur,

#                         **comparaison,
#                     }
#                 )

#         # ----------------------------------------------------
#         # 10. RESULTAT
#         # ----------------------------------------------------

#         print()
#         print("=" * 75)

#         if not erreurs:

#             print(
#                 "[OK] BUDGET VALIDE"
#             )

#             print(
#                 "Les résultats générés par SXMXDX "
#                 "correspondent à la Synthese Excel."
#             )

#             print("=" * 75)

#             return True

#         print(
#             "[ERREUR] BUDGET NON CONFORME"
#         )

#         print(
#             f"Nombre d'écarts : {len(erreurs)}"
#         )

#         print("=" * 75)

#         print()

#         print(
#             "DETAIL DES ECARTS"
#         )

#         print("-" * 75)

#         for erreur in erreurs:

#             print(
#                 f"{erreur['periode']:<12} "
#                 f"{erreur['indicateur']:<25} "
#                 f"Excel={afficher_nombre(erreur['attendu'])} "
#                 f"DB={afficher_nombre(erreur['obtenu'])} "
#                 f"Ecart={afficher_nombre(erreur['ecart'])}"
#             )

#         return False

#     finally:

#         if connexion is not None:

#             connexion.close()


# # ============================================================
# # MAIN
# # ============================================================

# if __name__ == "__main__":

#     try:

#         valide = valider_budget()

#         if valide:
#             sys.exit(0)

#         sys.exit(1)

#     except Exception as exc:

#         print()
#         print("=" * 75)
#         print(
#             "[ERREUR] VALIDATION IMPOSSIBLE"
#         )
#         print("=" * 75)

#         print(
#             str(exc)
#         )

#         sys.exit(1)







from decimal import Decimal, InvalidOperation
from pathlib import Path
import sys
import psycopg2
from psycopg2.extras import RealDictCursor
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

DB_CONFIG = {
    "host": "localhost",
    "port": 5432,
    "dbname": "sxmxdx2",
    "user": "sxmxdx",
    "password": "diary",
}

FEUILLE_SYNTHESE = "Synthese"
TOLERANCE_MONTANT = Decimal("1.00")
TOLERANCE_MARGE = Decimal("0.0001")
INDICATEURS = {
    "ca_export": ["ca", "export"],
    "charges_exploitation": ["charges", "exploitation"],
    "frais_export": ["frais", "export"],
    "masse_salariale": ["masse", "salariale"],
    "resultat_operationnel": ["resultat", "operationnel"],
    "marge_operationnelle": ["marge", "operationnelle"]
}

def connecter():
    return psycopg2.connect(**DB_CONFIG)

def decimaliser(valeur):
    if valeur is None:
        return Decimal("0")
    if isinstance(valeur, Decimal):
        return valeur
    try:
        return Decimal(str(valeur))
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError(f"Valeur numérique invalide : {valeur}")

def recuperer_dernier_budget(connexion):
    with connexion.cursor(cursor_factory=RealDictCursor) as curseur:
        curseur.execute("""
            SELECT b.id AS budget_id, b.import_id, b.exercice_id, b.ca_export,
                   b.charges_exploitation, b.frais_export, b.masse_salariale,
                   b.resultat_operationnel, b.marge_operationnelle, b.statut, e.annee
            FROM budget b
            JOIN exercice e ON e.id = b.exercice_id
            WHERE b.statut = 'GENERE' AND b.type_budget = 'GENERE'
            ORDER BY b.id DESC LIMIT 1
        """)
        budget = curseur.fetchone()
    if not budget:
        raise ValueError("Aucun budget GENERE trouvé.")
    return budget

def recuperer_import(connexion, import_id):
    with connexion.cursor(cursor_factory=RealDictCursor) as curseur:
        curseur.execute("SELECT * FROM import_fichier WHERE id = %s", (import_id,))
        ligne = curseur.fetchone()
    if not ligne:
        raise ValueError(f"Import {import_id} introuvable.")
    return ligne

def recuperer_budget_mensuel(connexion, budget_id):
    with connexion.cursor(cursor_factory=RealDictCursor) as curseur:
        curseur.execute("""
            SELECT bm.*, m.nom AS mois_nom
            FROM budget_mensuel bm
            JOIN mois m ON m.id = bm.mois_id
            WHERE bm.budget_id = %s
            ORDER BY bm.mois_id
        """, (budget_id,))
        lignes = curseur.fetchall()
    if len(lignes) != 12:
        raise ValueError(f"Le budget doit contenir exactement 12 mois. Trouvé : {len(lignes)}")
    return lignes

def trouver_fichier_excel(import_fichier):
    for colonne in ["chemin_fichier", "fichier", "nom_fichier", "filepath", "file_path"]:
        if import_fichier.get(colonne):
            chemin = Path(str(import_fichier[colonne]))
            if chemin.is_absolute() and chemin.exists():
                return chemin
            candidat = (PROJECT_ROOT / chemin).resolve()
            if candidat.exists():
                return candidat
            upload_dir = PROJECT_ROOT / "data" / "uploads"
            if upload_dir.exists():
                fichiers = sorted(upload_dir.rglob(chemin.name), key=lambda p: p.stat().st_mtime, reverse=True)
                if fichiers:
                    return fichiers[0]
    raise FileNotFoundError("Fichier Excel source introuvable.")

def normaliser_texte(valeur):
    return "" if valeur is None else str(valeur).strip().lower().translate(
        str.maketrans("éèêàôûï", "eee aoui".replace(" ", ""))
    )

def trouver_ligne_indicateur(feuille, mots_cles):
    mots_cles = [normaliser_texte(mot) for mot in mots_cles]
    for ligne in feuille.iter_rows():
        for cellule in ligne:
            texte = normaliser_texte(cellule.value)
            if texte and all(mot in texte for mot in mots_cles):
                return cellule.row
    return None

def extraire_ligne_synthese(feuille, ligne):
    return {
        "mensuel": [decimaliser(feuille.cell(row=ligne, column=c).value) for c in range(3, 15)],
        "annuel": decimaliser(feuille.cell(row=ligne, column=15).value)
    }

def extraire_synthese_excel(fichier_excel):
    classeur = load_workbook(fichier_excel, data_only=True, read_only=True)
    try:
        if FEUILLE_SYNTHESE not in classeur.sheetnames:
            raise ValueError(f"Feuille '{FEUILLE_SYNTHESE}' introuvable.")
        feuille = classeur[FEUILLE_SYNTHESE]
        resultat = {}
        for indicateur, mots_cles in INDICATEURS.items():
            ligne = trouver_ligne_indicateur(feuille, mots_cles)
            if ligne is None:
                raise ValueError(f"Indicateur introuvable dans Synthese : {indicateur}")
            resultat[indicateur] = extraire_ligne_synthese(feuille, ligne)
        return resultat
    finally:
        classeur.close()

def comparer(attendu, obtenu, tolerance):
    return abs(decimaliser(attendu) - decimaliser(obtenu)) <= tolerance

def valider_budget():
    connexion = connecter()
    try:
        budget = recuperer_dernier_budget(connexion)
        import_fichier = recuperer_import(connexion, budget["import_id"])
        fichier_excel = trouver_fichier_excel(import_fichier)
        budgets_mensuels = recuperer_budget_mensuel(connexion, budget["budget_id"])
        excel = extraire_synthese_excel(fichier_excel)

        for index, mois in enumerate(budgets_mensuels):
            for indicateur in INDICATEURS:
                tolerance = TOLERANCE_MARGE if indicateur == "marge_operationnelle" else TOLERANCE_MONTANT
                if not comparer(excel[indicateur]["mensuel"][index], mois[indicateur], tolerance):
                    return False

        for indicateur in INDICATEURS:
            tolerance = TOLERANCE_MARGE if indicateur == "marge_operationnelle" else TOLERANCE_MONTANT
            if not comparer(excel[indicateur]["annuel"], budget[indicateur], tolerance):
                return False

        return True
    finally:
        connexion.close()

if __name__ == "__main__":
    print(valider_budget())