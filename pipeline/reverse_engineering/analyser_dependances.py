from pathlib import Path
from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
import csv
import re
import sys


OUTPUT_DIR = Path("output/reverse_engineering")


def nettoyer_reference(valeur):
    """
    Ex:
    'Exportations'!$N$20 -> Exportations!N20
    """
    valeur = valeur.replace("$", "")
    valeur = valeur.replace("'", "")
    return valeur


def extraire_references(formule, feuille_courante):
    references = []

    try:
        tokens = Tokenizer(formule).items

        for token in tokens:

            if token.type == "OPERAND" and token.subtype == "RANGE":

                valeur = nettoyer_reference(token.value)

                # Ignore les constantes / noms non-cellulaires
                if "!" in valeur:
                    feuille_source, cellule_source = valeur.rsplit("!", 1)

                else:
                    feuille_source = feuille_courante
                    cellule_source = valeur

                # Vérifier que ça ressemble à une cellule/plage Excel
                if re.search(r"[A-Z]+\d+", cellule_source, re.IGNORECASE):

                    references.append({
                        "feuille_source": feuille_source,
                        "cellule_source": cellule_source
                    })

    except Exception as e:
        print(f"[ATTENTION] Formule non analysée : {formule} -> {e}")

    return references


def analyser_dependances(chemin_fichier):

    fichier = Path(chemin_fichier)

    if not fichier.exists():
        raise FileNotFoundError(fichier)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print("\n" + "=" * 75)
    print("SXMXDX - ANALYSE DES DEPENDANCES")
    print("=" * 75)

    wb = load_workbook(
        fichier,
        data_only=False
    )

    dependances = []

    for nom_feuille in wb.sheetnames:

        ws = wb[nom_feuille]

        compteur = 0

        for row in ws.iter_rows():

            for cell in row:

                formule = cell.value

                if not (
                    isinstance(formule, str)
                    and formule.startswith("=")
                ):
                    continue

                references = extraire_references(
                    formule,
                    nom_feuille
                )

                for ref in references:

                    dependances.append({
                        "feuille_destination": nom_feuille,
                        "cellule_destination": cell.coordinate,
                        "feuille_source": ref["feuille_source"],
                        "cellule_source": ref["cellule_source"],
                        "formule": formule
                    })

                    compteur += 1

        print(
            f"{nom_feuille:<25} : "
            f"{compteur} dépendance(s)"
        )

    wb.close()

    sortie = OUTPUT_DIR / "dependances.csv"

    with open(
        sortie,
        "w",
        newline="",
        encoding="utf-8"
    ) as f:

        champs = [
            "feuille_destination",
            "cellule_destination",
            "feuille_source",
            "cellule_source",
            "formule"
        ]

        writer = csv.DictWriter(
            f,
            fieldnames=champs
        )

        writer.writeheader()
        writer.writerows(dependances)

    print("\n" + "-" * 75)
    print(f"Total dépendances : {len(dependances)}")
    print(f"[OK] Généré : {sortie}")


if __name__ == "__main__":

    if len(sys.argv) != 2:
        print(
            "Utilisation : python analyser_dependances.py "
            "<fichier_excel>"
        )
        sys.exit(1)

    try:
        analyser_dependances(sys.argv[1])

    except Exception as e:
        print(f"[ERREUR] {e}")
        sys.exit(1)