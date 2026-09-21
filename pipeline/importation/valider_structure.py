from pathlib import Path
import sys

from openpyxl import load_workbook


FEUILLES_ATTENDUES = [
    "Donnees SOMIDA",
    "Employes",
    "Production",
    "Charges",
    "Exportations",
    "Frais Export",
    "Synthese",
]


def valider_structure(chemin_fichier):

    fichier = Path(chemin_fichier)

    print("\n" + "=" * 70)
    print("SXMXDX - VALIDATION DU FICHIER")
    print("=" * 70)

    if not fichier.exists():
        raise FileNotFoundError(
            f"Fichier introuvable : {fichier}"
        )

    if fichier.suffix.lower() not in [".xlsx", ".xlsm"]:
        raise ValueError(
            "Le fichier doit être au format XLSX ou XLSM."
        )

    print(f"[OK] Fichier : {fichier.name}")

    try:
        wb = load_workbook(
            fichier,
            read_only=True,
            data_only=False
        )

    except Exception as e:
        raise ValueError(
            f"Impossible de lire le classeur : {e}"
        )

    print("[OK] Classeur Excel lisible")

    feuilles = wb.sheetnames

    print("\nFeuilles détectées :")

    erreurs = []

    for feuille in FEUILLES_ATTENDUES:

        if feuille in feuilles:
            print(f"  [OK] {feuille}")

        else:
            print(f"  [MANQUANTE] {feuille}")
            erreurs.append(feuille)

    wb.close()

    print("\n" + "-" * 70)

    if erreurs:

        print("STATUT : FICHIER NON COMPATIBLE")
        print(
            "Feuilles manquantes : "
            + ", ".join(erreurs)
        )

        return False

    print("STATUT : FICHIER COMPATIBLE")

    return True


if __name__ == "__main__":

    if len(sys.argv) != 2:
        print(
            "Utilisation : python valider_structure.py "
            "<fichier_excel>"
        )
        sys.exit(1)

    try:

        valide = valider_structure(sys.argv[1])

        if not valide:
            sys.exit(1)

    except Exception as e:

        print(f"\n[ERREUR] {e}")
        sys.exit(1)